#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ساخت پاورپوینت فارسی و قابل ویرایش از معماری Prometheus."""

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

OUT = "prometheus-architecture-fa.pptx"
FONT = "Tahoma"

# ---------- palette ----------
NAVY      = RGBColor(0x1F, 0x38, 0x64)
NAVY_SOFT = RGBColor(0x2E, 0x4A, 0x7A)
INK       = RGBColor(0x1F, 0x2A, 0x37)
MUTED     = RGBColor(0x4B, 0x55, 0x63)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GREY_TXT  = RGBColor(0x6B, 0x72, 0x80)

BLUE   = (RGBColor(0xEA, 0xF2, 0xF9), RGBColor(0x5B, 0x8D, 0xB8))
GREEN  = (RGBColor(0xE6, 0xF4, 0xEC), RGBColor(0x4E, 0x9C, 0x6B))
PURPLE = (RGBColor(0xEF, 0xEC, 0xF9), RGBColor(0x7C, 0x6B, 0xB0))
ORANGE = (RGBColor(0xFD, 0xEE, 0xDF), RGBColor(0xD9, 0x8B, 0x4A))
GREY   = (RGBColor(0xF4, 0xF6, 0xF9), RGBColor(0xB0, 0xB8, 0xC4))
NAVYF  = (RGBColor(0xE7, 0xEC, 0xF5), RGBColor(0x1F, 0x38, 0x64))

# ---------- geometry ----------
SW, SH = 13.333, 7.5
LEFT, RIGHT = 0.45, 11.0
CW = RIGHT - LEFT                       # 10.55
BADGE_X, BADGE_W = 11.2, 1.68
GAP = 0.35
COL3 = (CW - 2 * GAP) / 3               # 3.2833
COL4 = (CW - 3 * GAP) / 4               # 2.375
X3 = [RIGHT - COL3 - i * (COL3 + GAP) for i in range(3)]     # RTL columns
C3 = [x + COL3 / 2 for x in X3]
X4 = [RIGHT - COL4 - i * (COL4 + GAP) for i in range(4)]
C4 = [x + COL4 / 2 for x in X4]


# ---------- helpers ----------
def style_run(run, size, bold=False, color=INK, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def add_par(tf, text, size, bold=False, color=INK, space_before=1, first=False, align="ctr"):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pPr.set("rtl", "1")
    p.alignment = {"ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT,
                   "l": PP_ALIGN.LEFT}[align]
    p.line_spacing = 1.0
    p.space_before = Pt(space_before)
    p.space_after = Pt(0)
    style_run(p.add_run(), size, bold, color)
    p.runs[0].text = text
    return p


def box(slide, x, y, w, h, title, lines=(), colors=BLUE, name="box",
        tsize=10.5, bsize=8.5, title_color=INK, body_color=MUTED,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, align="ctr", dashed=False):
    fill, line = colors
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = name
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(1.1)
    if dashed:
        sh.line.dash_style = 4
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, 0.085 / min(w, h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.07)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_par(tf, title, tsize, True, title_color, 0, True, align)
    for ln in lines:
        add_par(tf, ln, bsize, False, body_color, 1.5, False, align)
    return sh


def text(slide, x, y, w, h, value, size=10, bold=False, color=MUTED,
         align="ctr", name="text"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_par(tf, value, size, bold, color, 0, True, align)
    return tb


def arrow(slide, x1, y1, x2, y2, color=NAVY_SOFT, width=1.4, name="arrow"):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.name = name
    cn.line.color.rgb = color
    cn.line.width = Pt(width)
    ln = cn.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "sm")
    tail.set("len", "med")
    return cn


def header(slide, badge_no, badge_title, title, subtitle):
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(BADGE_X), Inches(0.25), Inches(BADGE_W), Inches(1.0))
    b.name = "phase_badge"
    b.shadow.inherit = False
    b.fill.solid()
    b.fill.fore_color.rgb = NAVY
    b.line.fill.background()
    b.adjustments[0] = 0.1
    tf = b.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_par(tf, badge_no, 13, True, WHITE, 0, True)
    add_par(tf, badge_title, 8.5, False, RGBColor(0xD8, 0xE0, 0xEE), 2)
    text(slide, LEFT, 0.24, CW, 0.62, title, 18, True, NAVY, "r", "slide_title")
    text(slide, LEFT, 0.86, CW, 0.42, subtitle, 10, False, GREY_TXT, "r", "slide_subtitle")


def plain_header(slide, title, subtitle=None):
    text(slide, LEFT, 0.42, CW, 0.7, title, 20, True, NAVY, "r", "slide_title")
    if subtitle:
        text(slide, LEFT, 1.02, CW, 0.36, subtitle, 10, False, GREY_TXT, "r", "slide_subtitle")


# ===========================================================================
def slide_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(LEFT), Inches(1.15), Inches(CW), Inches(0.06))
    band.name = "title_rule"
    band.shadow.inherit = False
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()

    text(s, LEFT, 1.45, CW, 0.5, "معماری سیستم — سه فاز", 12, True,
         RGBColor(0x5B, 0x8D, 0xB8), "r", "eyebrow")
    text(s, LEFT, 2.0, CW, 1.5,
         "پرامتئوس: پیش‌بینی فرآیند با درک تأخیر و سه جریان موازی",
         30, True, NAVY, "r", "main_title")
    text(s, LEFT, 3.45, CW, 0.6,
         "ساختار کامل مدل در سه فاز — نسخهٔ پاورپوینت، همهٔ متن‌ها و شکل‌ها قابل ویرایش",
         13, False, MUTED, "r")

    chips = [("فاز ۱", "مشخصات سیستم"), ("فاز ۲", "طراحی معماری"), ("فاز ۳", "طراحی جزئی")]
    for (no, nm), x in zip(chips, X3):
        c = box(s, x, 4.55, COL3, 1.05, no, [nm], NAVYF, f"chip_{no}",
                tsize=14, bsize=9.5, title_color=NAVY, body_color=NAVY)
        c.line.width = Pt(1.4)
    text(s, LEFT, 6.55, CW, 0.4,
         "بر پایهٔ main.py (مدل اصلی) و utils.py (پیاده‌سازی گراف عصبی پیشوند)",
         9, False, GREY_TXT, "r")
    return s


def slide_overview(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    plain_header(s, "نمای کلی: از داده خام تا پیش‌بینی",
                 "سه فاز پشت سر هم؛ هر فاز ورودی فاز بعد را می‌سازد")
    items = [
        ("فاز ۱", "مشخصات سیستم",
         ["تعریف مسئله و خروجی‌ها", "ورودی‌ها و آماده‌سازی داده", "ساخت ویژگی‌ها و گراف پیشوند"], BLUE),
        ("فاز ۲", "طراحی معماری",
         ["سه جریان موازی", "Transformer + TCN + GNN", "ادغام اطلاعات سه جریان"], GREEN),
        ("فاز ۳", "طراحی جزئی",
         ["لایه‌ها، ابعاد و پارامترها", "چهار سر خروجی", "راه‌اندازی و آموزش مدل"], PURPLE),
    ]
    for (no, nm, lines, col), x in zip(items, X3):
        box(s, x, 1.85, COL3, 3.0, no, [nm] + lines, col, f"phase_{no}",
            tsize=15, bsize=10.5, body_color=MUTED)
    for a, b in ((0, 1), (1, 2)):
        arrow(s, X3[a], 3.35, X3[b] + COL3, 3.35, NAVY, 2.0, "flow_arrow")
    text(s, LEFT, 5.25, CW, 0.45,
         "مسیر کار: داده ← ویژگی و گراف ← سه جریان موازی ← ادغام ← چهار خروجی",
         10.5, False, MUTED, "r")
    n = box(s, LEFT, 5.95, CW, 0.7,
            "همهٔ کادرها، فلش‌ها، رنگ‌ها و متن‌ها در پاورپوینت قابل ویرایش هستند.",
            colors=GREY, name="note_bar", tsize=10, dashed=True)
    return s


def slide_phase1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "فاز ۱", "مشخصات سیستم",
           "مشخصات سیستم: از لاگ رویداد تا وظایف یادگیری",
           "ورودی‌ها، ویژگی‌ها و اهداف مدل، به زبان ساده")

    y1, h1, y2, h2 = 1.5, 1.7, 3.85, 1.7
    top = [
        ("ورودی: لاگ رویداد (XES)",
         ["پرونده‌ها، فعالیت‌ها، زمان‌ها", "شناسهٔ هر پرونده", "ترتیب رخدادها"], BLUE),
        ("ساخت پیشوند",
         ["بریدن هر پرونده به پیشوند", "کمترین طول پیشوند: ۲", "پیشوند = نقطهٔ پیش‌بینی"], BLUE),
        ("آمار زمانی و موردی",
         ["زمان انتظار و مدت کل", "زمان باقی‌مانده", "نشانهٔ گلوگاه"], GREEN),
        ("اهداف پیش‌بینی",
         ["فعالیت بعدی", "زمان باقی‌مانده", "کاهش تأخیر و تشخیص گلوگاه"], ORANGE),
    ]
    for (t, ls, col), x in zip(top, X4):
        box(s, x, y1, COL4, h1, t, ls, col, f"p1_top_{t[:6]}", tsize=10.5, bsize=8.5)

    for a, b in ((0, 1), (1, 2), (2, 3)):
        arrow(s, X4[a], y1 + h1 / 2, X4[b] + COL4, y1 + h1 / 2, NAVY_SOFT, 1.6, "p1_chain")

    bottom = [
        ("ویژگی‌های ترتیبی و زمانی",
         ["شناسهٔ فعالیت و فعالیت قبلی", "موقعیت، فاصلهٔ زمانی، زمان تجمعی", "نقش و فراوانی فعالیت"], BLUE),
        ("گراف پیشوند",
         ["گره‌ها: فعالیت + ویژگی زمان و موقعیت", "یال‌ها: ترتیب و تکرار فعالیت‌ها", "روابط زمانی میان گره‌ها"], PURPLE),
        ("تنظیمات یادگیری",
         ["یادگیری چندوظیفه‌ای با نظارت", "تقسیم آموزش / اعتبارسنجی / آزمون", "بر اساس پرونده، بدون نشت داده"], GREEN),
    ]
    for (t, ls, col), x in zip(bottom, X4[1:]):
        box(s, x, y2, COL4, h2, t, ls, col, f"p1_bottom_{t[:6]}", tsize=10.5, bsize=8.5)

    for x in C4[1:]:
        arrow(s, x, y1 + h1, x, y2, NAVY_SOFT, 1.6, "p1_branch")

    box(s, X4[0], y2, COL4, h2,
        "نکته",
        ["هر پرونده به پیشوندهای کوتاه‌تر بریده می‌شود",
         "تا مدل پیش‌بینی زودهنگام را یاد بگیرد"],
        GREY, "p1_note", tsize=10, bsize=8.5, dashed=True)

    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(LEFT), Inches(6.0), Inches(CW), Inches(0.72))
    b.name = "p1_summary"
    b.shadow.inherit = False
    b.fill.solid()
    b.fill.fore_color.rgb = NAVY
    b.line.fill.background()
    b.adjustments[0] = 0.16
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_par(tf, "خروجی فاز ۱: دادهٔ آماده برای آموزش + تعریف روشن چهار وظیفهٔ پیش‌بینی",
            10.5, True, WHITE, 0, True)
    return s


def slide_phase2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "فاز ۲", "طراحی معماری",
           "معماری: سه جریان موازی و ادغام آن‌ها",
           "هر جریان یک دیدگاه متفاوت از پیشوند را می‌بیند")

    # shared input
    sx, sw_ = C3[1] - 2.8, 5.6
    box(s, sx, 1.5, sw_, 1.0, "بازنمایی ورودی مشترک",
        ["فعالیت + فعالیت قبلی + موقعیت",
         "کدگذاری زمان، نقش و فراوانی فعالیت"], NAVYF, "p2_shared",
        tsize=11, bsize=8.5, title_color=NAVY)

    ys, hs = 3.0, 1.7
    streams = [
        ("جریان ۱ — Transformer", ["کدگذار علّی", "یادگیری وابستگی‌های دوربرد", "ترتیب فعالیت‌ها"], BLUE),
        ("جریان ۲ — TCN", ["شبکهٔ پیچشی زمانی", "الگوهای محلی در چند مقیاس", "کرنل ۳ با گشودگی ۱، ۲، ۴"], GREEN),
        ("جریان ۳ — GNN گراف پیشوند", ["شبکهٔ عصبی گراف", "روابط بین فعالیت‌های پیشوند", "ورودی: گره‌ها و یال‌های گراف"], PURPLE),
    ]
    for (t, ls, col), x in zip(streams, X3):
        box(s, x, ys, COL3, hs, t, ls, col, f"p2_stream_{t[:10]}", tsize=11, bsize=8.5)

    # shared -> streams
    for x0, x1 in ((C3[1] + 0.35, C3[0]), (C3[1], C3[1]), (C3[1] - 0.35, C3[2])):
        arrow(s, x0, 2.5, x1, ys, NAVY_SOFT, 1.4, "p2_fan")

    # streams -> fusion
    yf, hf = 5.1, 1.35
    box(s, X3[1], yf, COL3 * 2 + 0.35, hf, "ادغام توجه متقابل (Transformer و TCN)",
        ["هر جریان با توجه به خروجی جریان دیگر بازبینی می‌شود",
         "توجه چندسری + اتصال باقیمانده"], ORANGE, "p2_cross", tsize=10.5, bsize=8.5)
    box(s, X3[2], yf, COL3, hf, "ادغام سه‌جریانه",
        ["ترتیب + گراف + تعامل", "ساخت بازنمایی نهایی"], ORANGE, "p2_tri",
        tsize=10.5, bsize=8.5)

    for x in (C3[0], C3[1]):
        arrow(s, x, ys + hs, x, yf, NAVY_SOFT, 1.6, "p2_to_cross")
    arrow(s, C3[2], ys + hs, C3[2], yf, NAVY_SOFT, 1.6, "p2_to_tri")
    arrow(s, X3[1], yf + hf / 2, X3[2] + COL3, yf + hf / 2, NAVY, 2.0, "p2_head")

    text(s, LEFT, 6.62, CW, 0.45,
         "سه جریان هم‌زمان و مستقل محاسبه می‌شوند؛ ادغام اول بین دو جریان و ادغام نهایی با جریان گراف انجام می‌شود.",
         9.5, False, GREY_TXT, "r")
    return s


def slide_phase3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "فاز ۳", "طراحی جزئی",
           "طراحی جزئی: لایه‌ها، ابعاد و سرهای خروجی",
           "ابعاد پایه: d_model = ۲۵۶")

    def row(y, h):
        return y, h

    # --- row A: input encodings
    yA, hA = 1.45, 0.85
    a = [
        ("کدگذاری ورودی (d_model = ۲۵۶)", ["تعبیهٔ فعالیت و فعالیت قبلی", "تعبیهٔ موقعیت و موقعیت نسبی"], BLUE),
        ("کدگذاری زمانی", ["لگاریتم فاصلهٔ زمانی و زمان تجمعی", "کدگذار نقش و فراوانی فعالیت"], BLUE),
        ("نرمال‌سازی و تنظیم", ["LayerNorm • Dropout ۰٫۰۵", "ماسک پدینگ و ماسک علّی"], GREY),
    ]
    for (t, ls, col), x in zip(a, X3):
        box(s, x, yA, COL3, hA, t, ls, col, f"p3_enc_{t[:8]}", tsize=10, bsize=8)

    # --- row B: encoders
    yB, hB = 2.6, 1.05
    b = [
        ("کدگذار Transformer", ["۶ لایه • ۸ سر توجه", "بُعد پیش‌خور ۱۰۲۴", "باقیمانده + LayerNorm"], BLUE),
        ("بلوک TCN", ["۱۰ پیچش • کرنل ۳", "گشودگی ۱، ۲، ۴", "باقیمانده + LayerNorm + GELU"], GREEN),
        ("گراف پیشوند (GNN)", ["تعبیهٔ فعالیت + ویژگی گره", "۲ لایهٔ GCN + LayerNorm + GELU", "تجمیع میانگین ماسک‌شده"], PURPLE),
    ]
    for (t, ls, col), x in zip(b, X3):
        box(s, x, yB, COL3, hB, t, ls, col, f"p3_enc2_{t[:8]}", tsize=10, bsize=8)

    for x in C3:
        arrow(s, x, yA + hA, x, yB, NAVY_SOFT, 1.4, "p3_to_enc")

    # --- row C: fusion
    yC, hC = 3.95, 0.9
    box(s, X3[1], yC, COL3 * 2 + 0.35, hC, "ادغام توجه متقابل",
        ["توجه چندسری + اتصال باقیمانده", "FFN + LayerNorm"], ORANGE, "p3_cross", tsize=10, bsize=8)
    box(s, X3[2], yC, COL3, hC, "ادغام ترتیبی",
        ["آخرین موقعیت معتبر پیشوند (پیش‌فرض)", "خروجی یک بردار برای هر پیشوند"], GREY, "p3_pool",
        tsize=10, bsize=8)
    for x in C3:
        arrow(s, x, yB + hB, x, yC, NAVY_SOFT, 1.4, "p3_to_fuse")

    # --- row D: tri-fusion MLP
    yD, hD = 5.15, 0.75
    box(s, LEFT, yD, CW, hD, "ادغام نهایی — MLP سه‌جریانه",
        ["الحاق [بازنمایی ترتیبی، گراف، تعامل] • ابعاد ۷۶۸ به ۵۱۲ به ۲۵۶  •  بازنمایی مشترک: LayerNorm + Linear + GELU"],
        ORANGE, "p3_mlp", tsize=10, bsize=8)
    arrow(s, C3[0] + 1.1, yC + hC, C3[0] + 1.1, yD, NAVY_SOFT, 1.4, "p3_to_mlp_r")
    arrow(s, C3[2], yC + hC, C3[2], yD, NAVY_SOFT, 1.4, "p3_to_mlp_l")

    # --- row E: heads
    yE, hE = 6.2, 0.85
    w4, gap = (CW - 3 * 0.25) / 4, 0.25
    xE = [RIGHT - w4, RIGHT - 2 * w4 - gap, RIGHT - 3 * w4 - 2 * gap, RIGHT - 4 * w4 - 3 * gap]
    cE = [x + w4 / 2 for x in xE]
    heads = [
        ("سر فعالیت", ["فعالیت بعدی", "خروجی دسته‌بندی"], BLUE),
        ("سر زمان باقی‌مانده", ["زمان باقی‌مانده", "رگرسیون"], GREEN),
        ("سر بهبود", ["کاهش تأخیر", "رگرسیون + Tanh"], GREEN),
        ("سر گلوگاه", ["گلوگاه هر موقعیت", "خروجی دسته‌بندی"], PURPLE),
    ]
    for (t, ls, col), x in zip(heads, xE):
        box(s, x, yE, w4, hE, t, ls, col, f"p3_head_{t[:8]}", tsize=10, bsize=8)
    for x in cE:
        arrow(s, x, yD + hD, x, yE, NAVY_SOFT, 1.4, "p3_to_head")

    text(s, LEFT, 7.12, CW, 0.3,
         "آموزش: زیان وزنی چندوظیفه‌ای (Cross-Entropy + Huber + Huber + BCE) • بهینه‌ساز AdamW • Warmup + Cosine • برش گرادیان • SWA",
         8, False, GREY_TXT, "r")
    return s


def main():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    slide_title(prs)
    slide_overview(prs)
    slide_phase1(prs)
    slide_phase2(prs)
    slide_phase3(prs)
    prs.save(OUT)
    print("saved", OUT, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
