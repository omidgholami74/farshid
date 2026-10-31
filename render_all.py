# -*- coding: utf-8 -*-
"""Build one editable PPTX per figure (plus a PIL preview PNG for checking)."""
import glob
import importlib
import os
import re
import sys

import figkit
from figkit import REG, Fig, bottom, cx, right
from pptx import Presentation
from pptx.util import Emu

for path in sorted(glob.glob("figdefs*.py")):
    importlib.import_module(path[:-3])


def num(name):
    m = re.search(r"3_(\d+)_", name)
    return int(m.group(1)) if m else 999


def check(fig):
    """Report shapes that fall outside the canvas."""
    W, H = fig.W, fig.H
    bad = []
    for sh in fig.slide.shapes:
        x0 = sh.left / 914400.0
        y0 = sh.top / 914400.0
        x1 = x0 + sh.width / 914400.0
        y1 = y0 + sh.height / 914400.0
        if x0 < -0.02 or y0 < -0.02 or x1 > W + 0.02 or y1 > H + 0.02:
            bad.append((sh.name, round(x0, 2), round(y0, 2),
                        round(x1, 2), round(y1, 2)))
    return bad


def main():
    out = "fig-pptx"
    prev = "artifacts/figs"
    os.makedirs(out, exist_ok=True)
    os.makedirs(prev, exist_ok=True)
    problems = []
    for name in sorted(REG, key=num):
        path = figkit.build(name, out, pngdir=prev)
        W, H, fn = REG[name]
        prs = Presentation(path)
        fig = Fig(prs.slides[0], W, H)
        bad = check(fig)
        if bad:
            problems.append((name, bad))
        print("%-34s %5.2f x %5.2f in  %2d shapes" %
              (name, W, H, len(prs.slides[0].shapes)))
    if problems:
        print("\nOUT OF BOUNDS:")
        for name, bad in problems:
            for b in bad:
                print("  ", name, b)
    else:
        print("\nall shapes inside the canvas")


if __name__ == "__main__":
    sys.exit(main())
