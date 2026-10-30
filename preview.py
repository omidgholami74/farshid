#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""پیش‌نمایش تقریبی + بررسی سرریز متن و هم‌پوشانی شکل‌ها."""

import sys
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

DPI = 110
REG = "fonts/Vazirmatn-Regular.ttf"
BLD = "fonts/Vazirmatn-Bold.ttf"
PPTX = sys.argv[1] if len(sys.argv) > 1 else "prometheus-architecture-fa.pptx"

_cache = {}


def font(size_px, bold):
    k = (round(size_px), bold)
    if k not in _cache:
        _cache[k] = ImageFont.truetype(BLD if bold else REG, max(6, round(size_px)))
    return _cache[k]


def shape_text(runs_shape):
    out = []
    for p in runs_shape.text_frame.paragraphs:
        txt = "".join(r.text for r in p.runs)
        if not txt.strip():
            continue
        r0 = p.runs[0]
        size = r0.font.size.pt if r0.font.size else 10
        bold = bool(r0.font.bold)
        col = (34, 48, 63)
        try:
            if r0.font.color and r0.font.color.rgb:
                c = r0.font.color.rgb
                col = (c[0], c[1], c[2])
        except Exception:
            pass
        align = str(p.alignment).split(".")[-1].split(" ")[0] if p.alignment else "CENTER"
        out.append((txt, size, bold, col, align, p.space_before.pt if p.space_before else 0))
    return out


def draw_slide(slide, path):
    W, H = 13.333 * DPI, 7.5 * DPI
    img = Image.new("RGB", (int(W), int(H)), "white")
    d = ImageDraw.Draw(img)
    problems = []
    rects = []

    for sh in slide.shapes:
        x, y = Emu(sh.left).inches * DPI, Emu(sh.top).inches * DPI
        w, h = Emu(sh.width).inches * DPI, Emu(sh.height).inches * DPI

        # bounds check
        if x < -1 or y < -1 or x + w > W + 1 or y + h > H + 1:
            problems.append(f"OUT OF SLIDE: {sh.name} ({x/DPI:.2f},{y/DPI:.2f},{w/DPI:.2f},{h/DPI:.2f})")

        if sh.shape_type == MSO_SHAPE_TYPE.LINE or sh.shape_type == 20 or sh.__class__.__name__ == "Connector":
            el = sh._element
            flipH = el.spPr.xfrm.get("flipH") in ("1", "true")
            flipV = el.spPr.xfrm.get("flipV") in ("1", "true")
            x1 = x + (w if flipH else 0)
            x2 = x + (0 if flipH else w)
            y1 = y + (h if flipV else 0)
            y2 = y + (0 if flipV else h)
            d.line([x1, y1, x2, y2], fill=(46, 74, 122), width=2)
            import math
            ang = math.atan2(y2 - y1, x2 - x1)
            L, aw = 9, 4.5
            d.polygon([(x2, y2),
                       (x2 - L * math.cos(ang) + aw * math.sin(ang), y2 - L * math.sin(ang) - aw * math.cos(ang)),
                       (x2 - L * math.cos(ang) - aw * math.sin(ang), y2 - L * math.sin(ang) + aw * math.cos(ang))],
                      fill=(46, 74, 122))
            continue

        # filled shape
        try:
            fill = tuple(sh.fill.fore_color.rgb)
        except Exception:
            fill = None
        try:
            line = tuple(sh.line.color.rgb) if sh.line.width and sh.line.width.pt > 0 else None
        except Exception:
            line = None
        if fill:
            rects.append((sh.name, x, y, w, h))
        if fill or line:
            rad = 12
            if sh.__class__.__name__ == "Shape":
                try:
                    rad = sh.adjustments[0] * min(w, h)
                except Exception:
                    pass
            d.rounded_rectangle([x, y, x + w, y + h], radius=rad, fill=fill,
                                outline=line or None, width=2)

        if not sh.has_text_frame:
            continue
        lines = shape_text(sh)
        if not lines:
            continue

        # measure
        rendered, total_h, max_over = [], 0, 0
        avail = w - 2 * 0.07 * DPI
        for (txt, size, bold, col, align, sb) in lines:
            disp = txt
            f = font(size * DPI / 72, bold)
            tw = f.getlength(disp)
            max_over = max(max_over, tw - avail)
            lh = size * DPI / 72 * 1.32
            total_h += lh + sb * DPI / 72
            rendered.append((disp, f, tw, lh, col, align, sb))

        if max_over > 2:
            problems.append(f"TEXT OVERFLOW ({max_over/DPI:.2f}in): {sh.name} -> {lines[0][0][:42]}")
        if total_h > h - 2 * 0.03 * DPI + 1:
            problems.append(f"TEXT TALL ({total_h/DPI:.2f}in > {h/DPI:.2f}in): {sh.name} -> {lines[0][0][:42]}")

        cy = y + h / 2 - total_h / 2 + 2
        for (disp, f, tw, lh, col, align, sb) in rendered:
            cy += sb * DPI / 72
            if align == "RIGHT":
                tx = x + w - 0.07 * DPI - tw
            elif align == "LEFT":
                tx = x + 0.07 * DPI
            else:
                tx = x + w / 2 - tw / 2
            d.text((tx, cy), disp, font=f, fill=col)
            cy += lh

    for i in range(len(rects)):
        n1, x1, y1, w1, h1 = rects[i]
        for j in range(i + 1, len(rects)):
            n2, x2, y2, w2, h2 = rects[j]
            ox = min(x1 + w1, x2 + w2) - max(x1, x2)
            oy = min(y1 + h1, y2 + h2) - max(y1, y2)
            if ox > 2 and oy > 2:
                problems.append(f"OVERLAP {ox/DPI:.2f}x{oy/DPI:.2f}in: {n1} / {n2}")

    img.save(path)
    return problems


prs = Presentation(PPTX)
all_problems = []
for i, slide in enumerate(prs.slides, 1):
    pr = draw_slide(slide, f"artifacts/slide{i}.png")
    all_problems += [f"slide{i}: {p}" for p in pr]

print("slides:", len(prs.slides))
print("\n".join(all_problems) if all_problems else "no layout problems detected")
