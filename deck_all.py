# -*- coding: utf-8 -*-
"""Collect every figure into a single review deck (one figure per slide)."""
import glob
import importlib
import re

from pptx import Presentation
from pptx.util import Emu

from figkit import REG, Fig

for path in sorted(glob.glob("figdefs*.py")):
    importlib.import_module(path[:-3])

SW, SH = 13.333, 7.9
MARGIN = 0.22
OUT = "figures-fa.pptx"


def num(name):
    m = re.search(r"3_(\d+)_", name)
    return int(m.group(1)) if m else 999


def main():
    prs = Presentation()
    prs.slide_width = Emu(int(SW * 914400))
    prs.slide_height = Emu(int(SH * 914400))
    blank = prs.slide_layouts[6]

    for name in sorted(REG, key=num):
        W, H, fn = REG[name]
        k = min((SW - 2 * MARGIN) / W, (SH - 2 * MARGIN) / H)
        dx = (SW - W * k) / 2.0
        dy = (SH - H * k) / 2.0

        slide = prs.slides.add_slide(blank)
        fig = Fig(slide, W, H, scale=k)
        fn(fig)
        for sh in slide.shapes:
            sh.left = Emu(int(sh.left + dx * 914400))
            sh.top = Emu(int(sh.top + dy * 914400))

        xs = [s.left / 914400.0 for s in slide.shapes]
        ys = [s.top / 914400.0 for s in slide.shapes]
        xe = [(s.left + s.width) / 914400.0 for s in slide.shapes]
        ye = [(s.top + s.height) / 914400.0 for s in slide.shapes]
        inside = (min(xs) >= -0.01 and min(ys) >= -0.01
                  and max(xe) <= SW + 0.01 and max(ye) <= SH + 0.01)
        print("%-32s scale %.2f  %2d shapes  %s" %
              (name, k, len(slide.shapes),
               "ok" if inside else "OUT OF SLIDE"))

    prs.save(OUT)
    print("\nsaved %s with %d slides" % (OUT, len(prs.slides._sldIdLst)))


if __name__ == "__main__":
    main()
