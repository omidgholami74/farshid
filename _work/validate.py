# -*- coding: utf-8 -*-
"""Validate the generated WordprocessingML: element ordering + structure."""
import os
import sys

import docx
from docx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_thesis as B  # noqa: E402

ORDER = {
    "w:pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
              "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
              "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
              "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
              "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
              "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
              "sectPr", "pPrChange"],
    "w:rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
              "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
              "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz",
              "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign",
              "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"],
    "w:tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
                "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
                "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption",
                "tblDescription", "tblPrChange"],
    "w:trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
               "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
               "ins", "del", "trPrChange"],
    "w:tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
               "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
               "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange"],
}


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else B.DOCX
    doc = docx.Document(path)
    print("file:", os.path.basename(path))
    body = doc.element.body
    problems = 0
    for tag, order in ORDER.items():
        for el in body.iter(qn(tag)):
            last = -1
            for ch in el:
                name = etree.QName(ch).localname
                if name not in order:
                    continue
                i = order.index(name)
                if i < last:
                    problems += 1
                    if problems <= 8:
                        print("ORDER  <%s>  child %r after %r" % (
                            etree.QName(el).localname, name, order[last]))
                last = max(last, i)

    direct = [tr for tr in body.iter(qn("w:tr")) if tr.find(qn("w:cantSplit")) is not None]
    print("ordering problems:", problems, "| cantSplit outside trPr:", len(direct))

    xml = doc.element.xml
    print("equations (m:oMath):", xml.count("<m:oMath"),
          "| tables:", len(doc.tables),
          "| updateFields:", "updateFields" in doc.settings.element.xml)
    eq = [t for t in doc.tables if len(t.columns) == 2 and len(t.rows) == 1]
    print("1x2 equation tables:", len(eq))
    nums = []
    for t in eq:
        ts = [x.text for x in t._tbl.iter(qn("w:t"))]
        nums.append(ts[0] if ts else "?")
    print("numbering:", nums[0], "...", nums[len(nums) // 2], "...", nums[-1])
    print("section present:", "معماری پایش عملیاتی مبتنی بر پرومتئوس" in xml)
    print("sources heading:", any(p.text.strip() == "منابع" for p in doc.paragraphs))
    return 1 if problems or direct else 0


if __name__ == "__main__":
    sys.exit(main())
