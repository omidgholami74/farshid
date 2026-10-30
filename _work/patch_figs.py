# -*- coding: utf-8 -*-
"""Insert the new Prometheus figure, justify body text, and tidy small defects."""
import copy
import os
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import patch_doc  # reuse the ordered-insert helper  # noqa: E402

DOCX = os.path.join(ROOT, "payannameh.docx")
PNG = os.path.join(HERE, "figs", "out", "prometheus.png")
CAPTION = ("شکل ۳.۲۳: معماری پایش عملیاتی مبتنی بر پرومتئوس و عامل هوشمند تصمیم\u200cیار")


def T(p):
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def main():
    doc = docx.Document(DOCX)
    paras = doc.paragraphs
    texts = [T(p) for p in paras]
    ref = next(i for i, t in enumerate(texts) if t.strip() == "منابع")

    # ---------- 1) justify body paragraphs ----------
    changed = 0
    for i in range(ref):
        p = paras[i]
        if p.style.name not in ("Normal", "ds-markdown-paragraph", "Paragraph"):
            continue
        t = texts[i].strip()
        if len(t) < 60 or t.startswith(("شکل", "جدول", "نمودار")):
            continue
        if p._p.findall(".//" + qn("w:drawing")) or p._p.findall(".//" + qn("w:pict")):
            continue
        ppr = p._p.get_or_add_pPr()
        jc = ppr.find(qn("w:jc"))
        cur = jc.get(qn("w:val")) if jc is not None else "inherit"
        if cur in ("center", "right", "both", "distribute"):
            continue
        if jc is None:
            jc = OxmlElement("w:jc")
            patch_doc.insert_ordered(ppr, jc)
        jc.set(qn("w:val"), "both")
        changed += 1

    doc.styles["Normal"].paragraph_format.alignment = AL.JUSTIFY

    # ---------- 2) Prometheus figure ----------
    existing = [t for t in texts if t.strip().startswith("شکل ۳.۲۳")]
    if existing:
        print("figure already present:", existing[0][:60])
    else:
        target = next(p for p in paras if T(p).strip().startswith("حالت عامل (Agent Mode)"))
        cap_src = next(p for p in paras if T(p).strip().startswith("شکل ۳.۲۲"))

        pic_p = target.insert_paragraph_before(style="Normal")
        pic_p.alignment = AL.CENTER
        pic_p.paragraph_format.space_before = Pt(6)
        pic_p.paragraph_format.space_after = Pt(4)
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "0")
        patch_doc.insert_ordered(pic_p._p.get_or_add_pPr(), bidi)
        pic_p.add_run().add_picture(PNG, width=Cm(15.6))

        cap_p = target.insert_paragraph_before(style=cap_src.style.name)
        cap_p.alignment = AL.CENTER
        cap_p.paragraph_format.space_before = Pt(2)
        cap_p.paragraph_format.space_after = Pt(10)
        run = cap_p.add_run(CAPTION)
        if cap_src.runs and cap_src.runs[0]._r.rPr is not None:
            run._r.insert(0, copy.deepcopy(cap_src.runs[0]._r.rPr))
        print("figure inserted before:", T(target).strip()[:50])

    # ---------- 3) tiny text fixes ----------
    fixes = 0
    for p in doc.paragraphs[:ref]:
        for t in p._p.iter(qn("w:t")):
            if t.text and "  " in t.text:
                new = t.text.replace("  ", " ")
                if new != t.text:
                    t.text = new
                    fixes += 1
    doc.save(DOCX)

    doc2 = docx.Document(DOCX)
    body = doc2.element.body
    print("justified:", changed, "| space fixes:", fixes)
    print("drawings:", len(body.findall(".//" + qn("w:drawing"))),
          "| equations:", len(body.findall(".//" + qn("m:oMath"))),
          "| tables:", len(doc2.tables),
          "| media:", len([r for r in doc2.part.rels.values() if "image" in r.reltype]))


if __name__ == "__main__":
    main()
