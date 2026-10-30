# -*- coding: utf-8 -*-
"""Swap the image of figure 3.23 for the freshly rendered PNG (keeps caption/style)."""
import os
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import patch_doc  # noqa: E402

DOCX = os.path.join(os.path.dirname(HERE), "payannameh.docx")
PNG = os.path.join(HERE, "figs", "out", "prometheus.png")


def T(p):
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def main():
    doc = docx.Document(DOCX)
    paras = doc.paragraphs
    idx = next(i for i, p in enumerate(paras) if T(p).strip().startswith("شکل ۳.۲۳"))
    cap = paras[idx]

    j = idx - 1
    while j >= 0 and not paras[j]._p.findall(".//" + qn("w:drawing")):
        j -= 1
    if j < 0:
        raise SystemExit("no picture paragraph before the caption")
    old = paras[j]
    old._p.getparent().remove(old._p)

    p = cap.insert_paragraph_before(style="Normal")
    p.alignment = AL.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "0")
    patch_doc.insert_ordered(p._p.get_or_add_pPr(), bidi)
    p.add_run().add_picture(PNG, width=Cm(15.6))

    doc.save(DOCX)

    chk = docx.Document(DOCX)
    inls = chk.element.body.findall(".//" + qn("wp:inline"))
    widest = max(inls, key=lambda e: int(e.find(qn("wp:extent")).get("cx"))) if inls else None
    if widest is not None:
        ext = widest.find(qn("wp:extent"))
        print("pictures:", len(inls),
              "| widest cm:", round(int(ext.get("cx")) / 360000, 2),
              "x", round(int(ext.get("cy")) / 360000, 2))


if __name__ == "__main__":
    main()
