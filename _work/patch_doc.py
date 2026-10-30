# -*- coding: utf-8 -*-
"""Apply the thesis patch (run from the project root)."""
import copy
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_thesis as B  # noqa: E402

W = B.w
M = B.m
fa = B.fa_num


# ------------------------------------------------------------------ utilities
def find_par(doc, prefix, required=True):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    if required:
        raise KeyError("anchor not found: %r" % prefix[:60])
    return None


def heading_para(ppr_template, text):
    """Heading paragraph that reuses an existing heading's pPr (numbering safe)."""
    p = B.el_from_string("<w:p xmlns:w='%s'/>" % B.W_NS)
    p.append(copy.deepcopy(ppr_template))
    p.append(B.make_run(text, rtl=True))
    return p


# ---------------------------------------------------------------------- main
def main():
    backup = os.path.join(HERE, "payannameh.original.docx")
    if not os.path.exists(backup):
        shutil.copy2(B.DOCX, backup)

    doc = docx.Document(B.DOCX)
    text_width = (doc.sections[0].page_width.twips
                  - doc.sections[0].left_margin.twips
                  - doc.sections[0].right_margin.twips)

    label = {}
    for i, k in enumerate(B.ORDER_CH3, 1):
        label[k] = "(%s-%s)" % (fa(3), fa(i))
    for i, k in enumerate(B.ORDER_CH4, 1):
        label[k] = "(%s-%s)" % (fa(4), fa(i))

    omml = B.build_omml_map(B.ORDER_CH3 + B.ORDER_CH4)
    inserted = 0

    # ------------------------------------------------ chapter 3: groups
    groups = []
    for key, anchor, intro in B.ANCHORS:
        if key is not None:
            groups.append({"anchor": anchor, "intro": intro, "keys": [key]})
        else:
            groups[-1]["keys"].append("NEXT")

    flat_keys = []
    for g in groups:
        keys = []
        for k in g["keys"]:
            if k == "NEXT":
                keys.append(B.ORDER_CH3[len(flat_keys)])
                flat_keys.append(keys[-1])
            else:
                keys.append(k)
                flat_keys.append(k)
        g["keys"] = keys
    if flat_keys != B.ORDER_CH3[:len(flat_keys)]:
        raise RuntimeError("anchor order does not match ORDER_CH3")

    for g in groups:
        base = find_par(doc, g["anchor"])._p
        labels = [label[k] for k in g["keys"]]
        intro = B.make_para(g["intro"].format(*labels), spacing_after=60)
        base.addnext(intro)
        base = intro
        for k in g["keys"]:
            tbl = B.make_eq_table(omml[k], label[k], text_width=text_width)
            base.addnext(tbl)
            base = tbl
            inserted += 1

    # ------------------------------------------------ chapter 4: groups
    ch4_groups = [
        {"anchor": "دقت در سطوح مختلف Top-k",
         "intro": ("معیار دقت در سطح k مطابق رابطه ({}) تعریف می‌شود؛ در این رابطه برچسب "
                   "واقعی نمونه i با مجموعه k پیش‌بینی برتر مدل مقایسه می‌گردد و نشانگر "
                   "براکتی مقدار یک را در صورت درست بودن شرط بازمی‌گرداند."),
         "keys": ["topk"]},
        {"anchor": "معیارهای Precision، Recall و F1 به تفکیک کلاس",
         "intro": ("معیارهای دقت، فراخوانی و F1 برای هر کلاس مطابق روابط ({})، ({}) و ({}) "
                   "تعریف می‌شوند. برای مقایسه کلی عملکرد مدل در شرایط عدم تعادل کلاس‌ها، "
                   "میانگین ماکروی F1 مطابق رابطه ({}) محاسبه می‌گردد."),
         "keys": ["precision", "recall", "f1", "macrof1"]},
    ]
    for g in ch4_groups:
        base = find_par(doc, g["anchor"])._p
        intro = B.make_para(g["intro"].format(*[label[k] for k in g["keys"]]),
                            spacing_after=60)
        base.addnext(intro)
        base = intro
        for k in g["keys"]:
            tbl = B.make_eq_table(omml[k], label[k], text_width=text_width)
            base.addnext(tbl)
            base = tbl
            inserted += 1

    # ------------------------------------------------ Prometheus + agent section
    if "معماری پایش عملیاتی مبتنی بر پرومتئوس" not in doc.element.xml:
        anchor_heading = find_par(doc, "جمع‌بندی مدل پیشنهادی")
        h2_ppr = anchor_heading._p.find(qn("w:pPr"))
        sample_h3 = find_par(doc, "خلاصه اجزای مدل و نوآوری‌های آن")
        h3_ppr = sample_h3._p.find(qn("w:pPr"))
        for el in build_section(label, omml, text_width, h2_ppr, h3_ppr):
            anchor_heading._p.addprevious(el)

    apply_text_fixes(doc)
    fix_captions(doc)
    fix_references(doc)
    enable_field_update(doc)

    doc.save(B.DOCX)
    print("equations inserted:", inserted)
    verify()
    return 0


# ------------------------------------------------------------------- section
def build_section(label, omml, text_width, h2_ppr, h3_ppr):
    blocks = []

    def P(text):
        blocks.append(("p", text))

    def H(level, text):
        blocks.append(("h%d" % level, text))

    def EQ(key):
        blocks.append(("eq", key))

    H(2, "معماری پایش عملیاتی مبتنی بر پرومتئوس و عامل هوشمند")
    P("مدل پیشنهادی در این فصل تا سطح یک سامانه پیش‌بینی تکمیل شده است، اما بهره‌برداری از "
      "آن در سازمان‌های مالی نیازمند لایه‌ای است که پیش‌بینی‌ها را پیوسته اندازه‌گیری، پایش و "
      "به اقدام عملیاتی تبدیل کند. در این بخش، معماری استقرار مدل بر پایه سامانه پایش "
      "پرومتئوس و یک عامل هوشمند توصیف می‌شود. این معماری مکمل مدل پیش‌بینی است و جایگزین آن "
      "محسوب نمی‌شود.")

    H(3, "ضرورت پایش عملیاتی و جایگاه آن در کاهش تأخیر")
    P("پایش پیش‌بینانه فرآیند به‌تنهایی تضمین‌کننده کاهش تأخیر نیست؛ پیش‌بینی زمانی به کاهش "
      "تأخیر منجر می‌شود که در یک چرخه عملیاتی به‌کار گرفته شود: اندازه‌گیری وضعیت جاری، "
      "مقایسه با هدف خدمت، هشدار پیش از وقوع انحراف و اقدام اصلاحی. در سامانه‌های مالی این "
      "چرخه باید با دوره تازگی کوتاه، در حد ثانیه تا دقیقه، اجرا شود، زیرا تأخیر در تشخیص، "
      "پنجره اقدام را از میان می‌برد.")
    P("شش الزام کلیدی برای این لایه قابل برشماری است: تازگی داده، توان پرس‌وجوی سریع روی "
      "سری‌های زمانی، هشدار قاعده‌مند و قابل ردیابی، مقیاس‌پذیری افقی، یکپارچگی با سامانه "
      "مدیریت نمونه فرآیندی، و تفکیک دسترسی و ممیزی اقدامات. معماری پرومتئوس به‌همراه یک عامل "
      "هوشمند، این الزامات را با کمترین پیچیدگی عملیاتی پوشش می‌دهد.")

    H(3, "معماری پرومتئوس و مؤلفه‌های آن")
    P("پرومتئوس یک سامانه متن‌باز پایش و پایگاه داده سری زمانی است که برخلاف سامانه‌های مبتنی "
      "بر ارسال، با مدل کشش کار می‌کند: سرور در بازه‌های منظم نقاط پایانی متریک را می‌خواند، "
      "نمونه‌ها را با برچسب‌های ابعادی در پایگاه داده سری زمانی ذخیره می‌کند و پرس‌وجوهای "
      "تحلیلی را با زبان PromQL پاسخ می‌دهد [15]. اجزای اصلی این معماری عبارت‌اند از سرور "
      "پرومتئوس، پایگاه داده سری زمانی، سازوکار کشف سرویس، صادرکننده‌ها، دروازه ارسال، مدیر "
      "هشدار و لایه نمایش.")
    P("جریان داده در این معماری یک‌طرفه و ساده است: صادرکننده‌ها متریک‌های دامنه و زیرساخت را "
      "در قالب متن ساده ارائه می‌کنند، سرور آن‌ها را در بازه scrape_interval می‌کشد، قواعد "
      "ثبت شاخص‌های تجمیعی را از پیش محاسبه می‌کنند، قواعد هشدار آستانه‌ها را می‌سنجند و مدیر "
      "هشدار، هشدارها را گروه‌بندی و به کانال‌های اطلاع‌رسانی مسیریابی می‌کند. مدل داده "
      "برچسب‌محور، امکان تفکیک متریک‌ها بر اساس نوع فرآیند، شعبه، محصول و مرحله فرآیندی را "
      "بدون تغییر در نام متریک فراهم می‌آورد و همین ویژگی، پرومتئوس را برای پایش چندفرآیندی "
      "مناسب می‌سازد.")
    P("در این پژوهش، هر نمونه فرآیندی در حال اجرا به یک سری زمانی با برچسب‌های case_type و "
      "stage نگاشت می‌شود؛ بنابراین شاخص‌هایی مانند زمان سپری‌شده نمونه، زمان انتظار در مرحله "
      "جاری و احتمال تأخیر پیش‌بینی‌شده در قالب متریک‌های استاندارد قابل پرس‌وجو خواهند بود. "
      "این نگاشت، شکاف میان خروجی مدل یادگیری عمیق و سامانه پایش عملیاتی را پر می‌کند.")

    H(3, "حالت عامل (Agent Mode) و صادرکننده‌های اختصاصی فرآیند")
    P("در استقرارهای چندشعبه‌ای، نمونه‌های پرومتئوس می‌توانند در حالت عامل اجرا شوند؛ در این "
      "حالت، نمونه تنها وظیفه کشش و ارسال داده به یک سرور مرکزی را بر عهده دارد و پرس‌وجو، "
      "قواعد و هشدار در سطح مرکزی انجام می‌شود [53]. این الگو برای سازمان‌های مالی که شعبات "
      "آن‌ها پشت شبکه‌های محدود قرار دارند و اجازه دسترسی ورودی به مرکز داده را ندارند مناسب "
      "است و بار محاسباتی و سطح حمله را کاهش می‌دهد.")
    P("صادرکننده اختصاصی فرآیند، مؤلفه‌ای است که از رویدادنگار سازمان یا صف رویداد تغذیه "
      "می‌شود و متریک‌های دامنه‌ای را در قالب استاندارد پرومتئوس ارائه می‌کند. در پیاده‌سازی "
      "پیشنهادی، این صادرکننده پنج متریک کلیدی را در معرض دید قرار می‌دهد: زمان سپری‌شده نمونه "
      "با نام ppm_case_elapsed_seconds، زمان باقی‌مانده پیش‌بینی‌شده با نام "
      "ppm_predicted_remaining_seconds، فعالیت بعدی پیش‌بینی‌شده به‌صورت شمارنده برچسب‌دار با "
      "نام ppm_predicted_next_activity_total، نشانگر گلوگاه با نام ppm_bottleneck_flag و "
      "اطمینان پیش‌بینی با نام ppm_prediction_confidence.")
    P("برای هموارسازی نویز لحظه‌ای، شاخص زمان چرخه به‌صورت میانگین متحرک نمایی محاسبه می‌شود؛ "
      "رابطه %s این هموارسازی را با ضریب α بیان می‌کند و مقدار کوچک α، وزن بیشتری به تاریخچه "
      "می‌دهد و از هشدارهای کاذب ناشی از نوسان‌های کوتاه‌مدت جلوگیری می‌نماید."
      % label["ewma"])
    EQ("ewma")

    H(3, "عامل هوشمند تصمیم‌یار و حلقه ادراک–استدلال–اقدام")
    P("عامل، مؤلفه نرم‌افزاری است که وضعیت فرآیند را ادراک می‌کند، درباره علت انحراف استدلال "
      "می‌کند و اقدام متناسب را انتخاب می‌کند. الگوی مورد استفاده در این معماری، الگوی استدلال "
      "و اقدام متناوب است که در آن عامل میان فراخوانی ابزارها و استدلال گام‌به‌گام تناوب ایجاد "
      "می‌کند [52]. ابزارهای در اختیار عامل شامل پرس‌وجوی PromQL، خدمت استنتاج مدل، سامانه "
      "مدیریت نمونه فرآیندی و سرویس اطلاع‌رسانی است.")
    P("حلقه عملکرد عامل چهار گام دارد: نخست ادراک، یعنی خواندن متریک‌ها، هشدارهای فعال و "
      "پیش‌بینی‌های جاری؛ دوم استدلال، یعنی مقایسه وضعیت با الگوی مرجع کشف‌شده از رویدادنگار، "
      "شناسایی گلوگاه و برآورد اثر تأخیر؛ سوم اقدام، یعنی اطلاع‌رسانی به مالک فرآیند، ایجاد "
      "وظیفه اصلاحی، اولویت‌دهی مجدد صف یا پیشنهاد مسیر جایگزین؛ و چهارم ثبت و بازخورد، یعنی "
      "ثبت اقدام، نتیجه و زمان پاسخ برای یادگیری بعدی.")
    P("انتخاب اقدام در هر گام بر پایه بیشینه سود مورد انتظار و با رعایت قیدهای عملیاتی انجام "
      "می‌شود (رابطه %s)؛ در این رابطه s_t وضعیت جاری، A مجموعه اقدام‌های مجاز و U تابع سود "
      "است که هزینه اقدام، اثر بر زمان چرخه و ریسک را ترکیب می‌کند. آستانه هشدار نیز مطابق "
      "رابطه %s تعریف می‌شود: اگر زمان باقی‌مانده پیش‌بینی‌شده از کسر معینی از هدف خدمت فراتر "
      "رود، هشدار فعال می‌گردد." % (label["policy"], label["alert"]))
    EQ("policy")
    EQ("alert")
    P("برای جلوگیری از اقدام‌های پرخطر، سه حفاظ در نظر گرفته می‌شود: نخست، حضور انسان در حلقه "
      "برای اقدام‌های با اثر مالی بالا؛ دوم، سقف نرخ اقدام خودکار در بازه زمانی مشخص؛ سوم، ثبت "
      "کامل زنجیره استدلال و ابزارهای فراخوانی‌شده برای ممیزی. عامل هرگز مجاز به تغییر مستقیم "
      "سوابق مالی نیست و اقدام‌های آن تنها به‌صورت پیشنهاد یا ایجاد وظیفه اعمال می‌شود.")

    H(3, "یکپارچه‌سازی با مدل پیش‌بینی و نمونه قواعد هشدار")
    P("مدل پیشنهادی از طریق یک خدمت استنتاج در دسترس قرار می‌گیرد و خروجی آن در هر گام به "
      "متریک تبدیل می‌شود. دو قاعده هشدار نمونه در این معماری تعریف شده است: نخست، اگر زمان "
      "باقی‌مانده پیش‌بینی‌شده برای یک نمونه بیش از دو برابر هدف خدمت شود و این وضعیت ده دقیقه "
      "ادامه یابد، هشدار سطح متوسط صادر می‌شود؛ دوم، اگر اطمینان پیش‌بینی کمتر از ۰٫۶ باشد، "
      "هشدار کیفیت داده برای بازبینی رویدادنگار ثبت می‌گردد. جدول ۳.۲ مؤلفه‌های این معماری و "
      "نقش هر یک را خلاصه می‌کند.")
    blocks.append(("caption", "جدول ۳.۲: مؤلفه‌های معماری پایش عملیاتی مبتنی بر پرومتئوس و "
                             "نقش هر مؤلفه"))
    blocks.append(("table", None))
    P("بازخورد اقدام‌ها به مدل نیز از همین مسیر تأمین می‌شود: نتیجه هر اقدام به‌صورت برچسب "
      "تأخیر واقعی در رویدادنگار ثبت می‌شود و در دوره‌های بعدی آموزش، به بهبود دقت پیش‌بینی "
      "زمان باقی‌مانده کمک می‌کند. این حلقه، پایش عملیاتی و یادگیری مدل را به یک چرخه پیوسته "
      "تبدیل می‌کند.")

    H(3, "ملاحظات پیاده‌سازی، امنیت و مقیاس‌پذیری")
    P("پیکربندی پیشنهادی برای استقرار، بازه کشش پانزده ثانیه، نگهداشت پانزده روزه داده خام و "
      "انتقال نمونه‌های تجمیعی به ذخیره‌سازی بلندمدت است؛ قواعد ثبت، شاخص‌های پرکاربرد را از "
      "پیش محاسبه می‌کنند تا هزینه پرس‌وجوهای داشبورد کاهش یابد. محدودیت شناخته‌شده پرومتئوس، "
      "تمرکز آن بر متریک‌های عددی است؛ بنابراین داده‌های رویدادی، متن استدلال عامل و سوابق "
      "اقدام در پایگاه داده عملیاتی نگهداری می‌شوند و تنها شاخص‌های عددی حاصل از آن‌ها به لایه "
      "پایش منتقل می‌گردد.")
    P("از منظر امنیت، ارتباط کشش و ارسال باید با TLS و احراز هویت انجام شود، دسترسی به رابط "
      "پرس‌وجو از طریق واسط احراز هویت‌کننده محدود گردد و داده‌های شخصی هرگز در برچسب‌ها قرار "
      "نگیرد، زیرا برچسب‌ها کاردینالیتی بالا و ریسک افشای اطلاعات ایجاد می‌کنند. از منظر "
      "مقیاس‌پذیری، نگهداشت کوتاه‌مدت داده خام همراه با نمونه‌برداری در فرآیندهای پرحجم و "
      "استفاده از ذخیره‌سازی بلندمدت برای تحلیل‌های تاریخی پیشنهاد می‌شود.")
    P("ارزیابی این لایه با دو معیار عملیاتی انجام می‌شود: میانگین زمان تشخیص تأخیر از لحظه "
      "وقوع انحراف تا صدور هشدار، و نرخ تحقق هدف خدمت در نمونه‌های تحت پایش. کاهش زمان تشخیص "
      "و افزایش نرخ تحقق هدف خدمت، ارزش عملیاتی معماری پیشنهادی را در کنار دقت پیش‌بینی مدل "
      "نشان می‌دهد.")

    rows = [
        ["مؤلفه", "نقش در معماری", "نمونه پیاده‌سازی"],
        ["سرور پرومتئوس", "کشش، ذخیره و پرس‌وجوی متریک‌ها با PromQL", "Prometheus Server 3.x"],
        ["پایگاه داده سری زمانی", "ذخیره فشرده نمونه‌ها با نگهداشت کوتاه‌مدت",
         "TSDB داخلی و ذخیره‌سازی بلندمدت"],
        ["صادرکننده فرآیند", "تبدیل رویدادهای فرآیندی به متریک‌های دامنه‌ای",
         "صادرکننده اختصاصی پایتون"],
        ["صادرکننده زیرساخت", "پایش منابع اجرای مدل و خدمت استنتاج",
         "Node Exporter و cAdvisor"],
        ["قواعد ثبت و هشدار", "پیش‌محاسبه شاخص‌ها و تعریف آستانه‌های هشدار",
         "Recording Rules و Alerting Rules"],
        ["مدیر هشدار", "گروه‌بندی، مسیریابی و مهار هشدارها", "Alertmanager"],
        ["عامل هوشمند", "ادراک، استدلال و اقدام خودکار با ابزارها",
         "عامل مبتنی بر LLM با الگوی ReAct"],
        ["داشبورد و گزارش", "نمایش وضعیت، روند و تحلیل ریشه‌ای", "Grafana"],
    ]

    out = []
    for kind, value in blocks:
        if kind == "h2":
            out.append(heading_para(h2_ppr, value))
        elif kind == "h3":
            out.append(heading_para(h3_ppr, value))
        elif kind == "p":
            out.append(B.make_para(value, spacing_after=140))
        elif kind == "eq":
            out.append(B.make_eq_table(omml[value], label[value], text_width=text_width))
        elif kind == "caption":
            out.append(B.make_para(value, jc="center", spacing_after=60))
        elif kind == "table":
            out.append(B.make_data_table(rows, [2000, 4000, 2503], text_width=text_width))
    out.append(B.make_para("", spacing_after=60))
    return out


# ------------------------------------------------------------------ text fixes
ZWNJ_FIXES = [
    ("تبدیل کننده", "تبدیل‌کننده"),
    ("شبکه های", "شبکه‌های"),
    ("می شود", "می‌شود"),
    ("می کند", "می‌کند"),
    ("می دهد", "می‌دهد"),
    ("داده های", "داده‌های"),
    ("مدل های", "مدل‌های"),
    ("فرآیند های", "فرآیندهای"),
    ("روش های", "روش‌های"),
    ("ویژگی های", "ویژگی‌های"),
    ("لایه های", "لایه‌های"),
    ("نمونه های", "نمونه‌های"),
    ("سامانه های", "سامانه‌های"),
    ("شاخص های", "شاخص‌های"),
]

COVER_FIXES = {
    "دانشگاه صنعتي امیرکبیر": "دانشگاه صنعتی امیرکبیر",
    "دانشكده مدیریت علم و فناوری": "دانشکده مدیریت علم و فناوری",
    "پایان‌نامه کارشناسی یا کارشناسی‌ارشد یا رسالۀ دکتری": "پایان‌نامه کارشناسی ارشد",
    "May & 2026": "May 2026",
    "اعضاي كميته": "اعضای کمیته",
}

DROP_PREFIXES = [
    "در صفحه تعهدنامه اصالت اثر، در قسمت بالا سمت چپ",
    "همچنین در صفحه تعهدنامه اصالت اثر، در خط اول",
]


def apply_text_fixes(doc):
    counts = {"zwnj": 0, "cover": 0, "dropped": 0}
    for p in list(doc.paragraphs):
        t = p.text.strip()
        if any(t.startswith(x) for x in DROP_PREFIXES):
            p._p.getparent().remove(p._p)
            counts["dropped"] += 1
            continue
        for run in p.runs:
            if not run.text:
                continue
            new = run.text
            for a, b in ZWNJ_FIXES:
                if a in new:
                    counts["zwnj"] += new.count(a)
                    new = new.replace(a, b)
            for a, b in COVER_FIXES.items():
                if a in new:
                    counts["cover"] += 1
                    new = new.replace(a, b)
            if new != run.text:
                run.text = new
    print("text fixes:", counts)


def _set_caption_style(p):
    ppr = p._p.get_or_add_pPr()
    for old in ppr.findall(qn("w:pStyle")):
        ppr.remove(old)
    ps = etree.Element(qn("w:pStyle"))
    ps.set(qn("w:val"), "shkl")
    ppr.insert(0, ps)


def fix_captions(doc):
    for p in doc.paragraphs:
        if p.text.strip().startswith("شکل ۴.۹:"):
            _set_caption_style(p)
            print("caption style fixed: شکل ۴.۹")
            break

    tbl = None
    for t in doc.tables:
        if t.rows and t.rows[0].cells and "مطالعه" in t.rows[0].cells[0].text:
            tbl = t
            break
    if tbl is None:
        return
    nxt = tbl._tbl.getnext()
    hops = 0
    while nxt is not None and hops < 3:
        if nxt.tag == qn("w:p"):
            txt = "".join(nxt.itertext()).strip()
            if txt.startswith("جدول ۴.۴:"):
                nxt.getparent().remove(nxt)
                tbl._tbl.addprevious(nxt)
                print("table 4.4 caption moved above the table")
                break
        nxt = nxt.getnext()
        hops += 1


def fix_references(doc):
    ref_par = None
    for p in doc.paragraphs:
        if p.text.strip() == "References":
            ref_par = p
            break
    if ref_par is not None:
        new_p = B.make_para("منابع", style_id=doc.styles["Heading 1"].style_id,
                            page_break_before=True, jc=None)
        ref_par._p.addprevious(new_p)
        ref_par._p.getparent().remove(ref_par._p)
        print("references heading replaced with منابع")

    for p in doc.paragraphs:
        if p.text.strip().startswith("[15] nbvmghm"):
            for run in p.runs:
                run.text = ""
            if p.runs:
                p.runs[0].text = (
                    "[15] Prometheus Authors, “Prometheus: Architecture overview,” Prometheus "
                    "Documentation, 2026. [Online]. Available: "
                    "https://prometheus.io/docs/introduction/architecture/")
            print("reference [15] replaced")
            break

    last = find_par(doc, "[51]", required=False)
    if last is not None:
        extra = [
            ("[52] S. Yao, J. Zhao, D. Yu, N. Du, I. Shafran, K. Narasimhan, and Y. Cao, "
             "“ReAct: Synergizing reasoning and acting in language models,” in Proc. Int. "
             "Conf. on Learning Representations (ICLR), 2023."),
            ("[53] Prometheus Authors, “Agent mode,” Prometheus Documentation, 2026. "
             "[Online]. Available: "
             "https://prometheus.io/docs/prometheus/latest/feature_flags/"),
        ]
        base = last._p
        for text in extra:
            p = B.make_para(text, jc="left", rtl=False, spacing_after=60)
            base.addnext(p)
            base = p
        print("references [52] and [53] appended")

    fixed = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if re.match(r"^\[\d+\]", txt):
            ppr = p._p.get_or_add_pPr()
            for tag in ("w:bidi", "w:jc"):
                for old in ppr.findall(qn(tag)):
                    ppr.remove(old)
            bidi = etree.Element(qn("w:bidi"))
            bidi.set(qn("w:val"), "0")
            jc = etree.Element(qn("w:jc"))
            jc.set(qn("w:val"), "left")
            insert_ordered(ppr, bidi)
            insert_ordered(ppr, jc)
            for run in p.runs:
                rpr = run._r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = etree.Element(qn("w:rPr"))
                    run._r.insert(0, rpr)
                for old in rpr.findall(qn("w:rtl")):
                    rpr.remove(old)
                rtl = etree.Element(qn("w:rtl"))
                rtl.set(qn("w:val"), "0")
                insert_ordered_rpr(rpr, rtl)
            fixed += 1
    print("reference paragraphs aligned LTR:", fixed)


PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
             "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
             "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
             "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
             "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
             "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]


RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
             "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
             "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz",
             "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign",
             "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]


def _insert_ordered(order, parent, element):
    idx = order.index(etree.QName(element).localname)
    for child in parent:
        name = etree.QName(child).localname
        if name in order and order.index(name) > idx:
            child.addprevious(element)
            return
    parent.append(element)


def insert_ordered(ppr, element):
    """Insert an element into pPr at its schema position."""
    _insert_ordered(PPR_ORDER, ppr, element)


def insert_ordered_rpr(rpr, element):
    """Insert an element into rPr at its schema position."""
    _insert_ordered(RPR_ORDER, rpr, element)


def enable_field_update(doc):
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        el = etree.SubElement(settings, qn("w:updateFields"))
        el.set(qn("w:val"), "true")
        print("updateFields enabled")


def verify():
    doc = docx.Document(B.DOCX)
    xml = doc.element.xml
    print("oMath count:", xml.count("<m:oMath"))
    print("tables:", len(doc.tables))
    print("Prometheus section:", "معماری پایش عملیاتی مبتنی بر پرومتئوس" in xml)
    eq_tables = [t for t in doc.tables if len(t.columns) == 2 and len(t.rows) == 1]
    print("equation tables (1x2):", len(eq_tables))
    t = eq_tables[0]
    tblPr = t._tbl.tblPr
    print("bidiVisual:", tblPr.find(qn("w:bidiVisual")) is not None,
          "| all borders none:",
          all(b.get(qn("w:val")) == "none"
              for b in tblPr.find(qn("w:tblBorders"))))


if __name__ == "__main__":
    sys.exit(main())
