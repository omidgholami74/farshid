# -*- coding: utf-8 -*-
"""Full diagnostic scan of the thesis: alignment, Latin text, text-formulas, typography."""
import os
import re
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(os.path.dirname(HERE), "payannameh.docx")
OUT = os.path.join(HERE, "scan.txt")

LATIN = re.compile(r"[A-Za-z]{2,}")
MATHY = re.compile(r"(?<![A-Za-z])(?:[A-Za-z]\s*=|∑|≤|≥|≈|×|√|\^|O\()")
LATIN_DIGIT = re.compile(r"[0-9]")
BAD_SPACE = [("  ", "double space"), (" ،", "space before comma"), (" .", "space before period"),
             (" )", "space before paren"), ("( ", "space after paren"), (" ؛", "space before ;"),
             (" :", "space before colon")]


def ptext(p):
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def jc_of(p):
    ppr = p._p.pPr
    if ppr is None:
        return "(inherit)"
    el = ppr.find(qn("w:jc"))
    return el.get(qn("w:val")) if el is not None else "(inherit)"


def main():
    doc = docx.Document(DOCX)
    paras = doc.paragraphs
    texts = [ptext(p) for p in paras]

    try:
        ref_start = next(i for i, t in enumerate(texts) if t.strip() == "منابع")
    except StopIteration:
        ref_start = len(texts)

    lines = []
    add = lines.append
    add("paragraphs: %d | tables: %d | refs start at: %d" % (len(paras), len(doc.tables), ref_start))

    # ---------- alignment ----------
    add("\n===== ALIGNMENT (paragraphs with text) =====")
    tally = {}
    for i, p in enumerate(paras):
        if not texts[i].strip():
            continue
        key = (p.style.name, jc_of(p))
        tally[key] = tally.get(key, 0) + 1
    for (style, jc), n in sorted(tally.items(), key=lambda kv: -kv[1]):
        add("  %-12s jc=%-10s %d" % (style, jc, n))

    add("\n--- body paragraphs NOT justified (style Normal, before منابع) ---")
    bad = 0
    for i in range(ref_start):
        p = paras[i]
        if not texts[i].strip() or p.style.name != "Normal":
            continue
        if jc_of(p) not in ("both", "distribute"):
            bad += 1
            if bad <= 60:
                add("%4d jc=%-9s %s" % (i, jc_of(p), texts[i].strip()[:110]))
    add("-- body paragraphs needing justify: %d" % bad)

    # ---------- headings ----------
    add("\n===== HEADINGS =====")
    for i, p in enumerate(paras):
        if p.style.name.startswith("Heading") and texts[i].strip():
            add("%4d  %-10s %s" % (i, p.style.name, texts[i].strip()[:90]))

    # ---------- latin text in body ----------
    add("\n===== LATIN IN BODY =====")
    latin = []
    for i in range(ref_start):
        t = texts[i].strip()
        if LATIN.search(t):
            latin.append(i)
            add("%4d [%s] %s" % (i, paras[i].style.name, t[:300].replace("\n", " ")))
    add("-- total latin paragraphs: %d" % len(latin))

    # ---------- latin in tables ----------
    add("\n===== LATIN IN TABLES =====")
    for ti, tbl in enumerate(doc.tables):
        for ri, row in enumerate(tbl.rows):
            joined = " | ".join(c.text.strip() for c in row.cells)
            if LATIN.search(joined) and len(joined) < 400:
                add("T%02d r%d: %s" % (ti, ri, joined[:200].replace("\n", " ")))

    # ---------- text formulas ----------
    add("\n===== TEXT-FORMULA CANDIDATES =====")
    for i in range(ref_start):
        t = texts[i].strip()
        if MATHY.search(t) and LATIN.search(t):
            add("%4d %s" % (i, t[:300].replace("\n", " ")))

    # ---------- captions ----------
    add("\n===== CAPTIONS =====")
    for i, t in enumerate(texts):
        s = t.strip()
        if re.match(r"^(شکل|جدول|نمودار|فرمول|رابطه)\s*[۰-۹0-9]", s) and len(s) < 200:
            add("%4d [%s] %s" % (i, paras[i].style.name, s[:150]))

    # ---------- typography ----------
    add("\n===== TYPOGRAPHY =====")
    for i in range(len(texts)):
        t = texts[i]
        if not t.strip():
            continue
        hits = [name for pat, name in BAD_SPACE if pat in t]
        if LATIN_DIGIT.search(t) and not re.match(r"^\s*\[\d+\]", t.strip()):
            hits.append("latin digits")
        if "\u200f" in t or "\u200e" in t:
            hits.append("bidi marks")
        if hits:
            add("%4d %-32s %s" % (i, ",".join(sorted(set(hits))), t[:150].replace("\n", " ")))

    # ---------- citations ----------
    add("\n===== CITATIONS =====")
    body = "\n".join(texts[:ref_start])
    cited = set()
    for m in re.finditer(r"\[([0-9,\s\-–]+)\]", body):
        for part in re.split(r"[,،]", m.group(1)):
            part = part.strip()
            if "-" in part or "–" in part:
                a, b = re.split(r"[-–]", part)
                if a.strip().isdigit() and b.strip().isdigit():
                    cited.update(range(int(a), int(b) + 1))
            elif part.isdigit():
                cited.add(int(part))
    refs = {}
    for i in range(ref_start, len(texts)):
        m = re.match(r"^\s*\[(\d+)\]", texts[i])
        if m:
            refs[int(m.group(1))] = texts[i].strip()
    add("cited: %d | listed: %d" % (len(cited), len(refs)))
    add("cited but not listed: %s" % sorted(cited - set(refs)))
    add("listed but never cited: %s" % sorted(set(refs) - cited))

    # ---------- misc ----------
    add("\n===== MISC =====")
    add("empty paragraphs: %d" % sum(1 for t in texts if not t.strip()))
    add("sections: %d" % len(doc.sections))
    for si, s in enumerate(doc.sections):
        add("  section %d: %dx%d twips, margins L%d R%d T%d B%d" % (
            si, s.page_width.twips, s.page_height.twips, s.left_margin.twips,
            s.right_margin.twips, s.top_margin.twips, s.bottom_margin.twips))
    add("inline OMML outside tables: %d" % len([
        m for m in doc.element.body.iter(qn("m:oMath"))
        if not any(a.tag == qn("w:tbl") for a in m.iterancestors())]))
    add("drawings: %d" % len(list(doc.element.body.iter(qn("w:drawing")))))
    for i, t in enumerate(texts):
        if re.search(r"(TODO|XXX|لطفاً|اینجا|placeholder)", t):
            add("PLACEHOLDER %4d: %s" % (i, t[:160]))

    report = "\n".join(lines)
    with open(OUT, "w", encoding="utf-8") as fh:
        fh.write(report)
    print("\n".join(lines[:200]))
    print("\n[full report: %s, %d lines]" % (OUT, len(lines)))


if __name__ == "__main__":
    main()
