# -*- coding: utf-8 -*-
"""Native (editable) PowerPoint charts with the same look as the figures."""
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION,
                             XL_LEGEND_POSITION, XL_MARKER_STYLE)
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from pptx.util import Pt

from figkit import FONT, LINE, MUTED, _rgb


def _patch_chart_font(chart, fs):
    chart.font.size = Pt(fs)
    chart.font.name = FONT
    txPr = chart._chartSpace.find(qn("c:txPr"))
    if txPr is None:
        return
    d = txPr.find(qn("a:defRPr"))
    if d is None:
        return
    for tag in ("a:latin", "a:ea", "a:cs"):
        for old in d.findall(qn(tag)):
            d.remove(old)
    for tag in ("a:latin", "a:ea", "a:cs"):
        d.append(d.makeelement(qn(tag), {"typeface": FONT}))


def _axis_title(ax, text, fs, color=MUTED):
    ax.has_title = True
    tf = ax.axis_title.text_frame
    tf.text = text
    for r in tf.paragraphs[0].runs:
        r.font.size = Pt(fs)
        r.font.name = FONT
        r.font.color.rgb = _rgb(color)


def bars(fig, x, y, w, h, cats, vals, color="#9DC3E6", ec=LINE, fs=9,
         xtitle=None, ytitle=None, ymax=None, gap=12, data_labels=False,
         horizontal=False, labels_fmt=None, per_point=None):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("فراوانی", vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = fig.slide.shapes.add_chart(kind, fig._emu(x), fig._emu(y),
                                    fig._emu(w), fig._emu(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    _patch_chart_font(ch, fs)

    plot = ch.plots[0]
    plot.gap_width = gap
    plot.vary_by_categories = False
    if data_labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(fs)
        dl.font.name = FONT
        if labels_fmt:
            dl.number_format = labels_fmt
            dl.number_format_is_linked = False

    ser = ch.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = _rgb(color)
    ser.format.line.color.rgb = _rgb(ec)
    ser.format.line.width = Pt(0.5)
    if per_point:
        for i, c in enumerate(per_point):
            if c and i < len(ser.points):
                pt = ser.points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = _rgb(c)

    ca = ch.category_axis
    ca.format.line.color.rgb = _rgb(MUTED)
    ca.tick_labels.font.size = Pt(fs)
    ca.tick_labels.font.name = FONT
    ca.has_major_gridlines = False

    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _rgb("#E0E0E0")
    va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.color.rgb = _rgb(MUTED)
    va.tick_labels.font.size = Pt(fs)
    va.tick_labels.font.name = FONT
    if ymax:
        va.maximum_scale = ymax
        va.minimum_scale = 0
    if ytitle and not horizontal:
        _axis_title(va, ytitle, fs + 1)
    if ytitle and horizontal:
        _axis_title(ca, ytitle, fs + 1)
    if xtitle and not horizontal:
        _axis_title(ca, xtitle, fs + 1)
    if xtitle and horizontal:
        _axis_title(va, xtitle, fs + 1)
    return ch


def lines(fig, x, y, w, h, cats, series, fs=9, ymin=None, ymax=None,
          xtitle=None, ytitle=None, legend=True):
    cd = CategoryChartData()
    cd.categories = cats
    for s in series:
        cd.add_series(s["name"], s["vals"])
    gf = fig.slide.shapes.add_chart(XL_CHART_TYPE.LINE, fig._emu(x),
                                    fig._emu(y), fig._emu(w), fig._emu(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = bool(legend)
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    _patch_chart_font(ch, fs)

    for s, ser in zip(series, ch.series):
        ser.format.line.color.rgb = _rgb(s.get("color", LINE))
        ser.format.line.width = Pt(s.get("lw", 2.0))
        ser.smooth = s.get("smooth", False)
        if s.get("dash"):
            ser.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        ser.marker.style = XL_MARKER_STYLE.NONE

    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(fs)
    ca.tick_labels.font.name = FONT
    ca.format.line.color.rgb = _rgb(MUTED)
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _rgb("#E0E0E0")
    va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.color.rgb = _rgb(MUTED)
    va.tick_labels.font.size = Pt(fs)
    va.tick_labels.font.name = FONT
    if ymin is not None:
        va.minimum_scale = ymin
    if ymax is not None:
        va.maximum_scale = ymax
    if ytitle:
        _axis_title(va, ytitle, fs + 1)
    if xtitle:
        _axis_title(ca, xtitle, fs + 1)
    return ch


def xy(fig, x, y, w, h, xvals, yvals, color="#1F3864", lw=2.2, fs=9,
       xtitle=None, ytitle=None):
    cd = XyChartData()
    ser_data = cd.add_series("Series 1")
    for xv, yv in zip(xvals, yvals):
        ser_data.add_data_point(xv, yv)
    gf = fig.slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
                                    fig._emu(x), fig._emu(y), fig._emu(w),
                                    fig._emu(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    _patch_chart_font(ch, fs)
    ser = ch.series[0]
    ser.format.line.color.rgb = _rgb(color)
    ser.format.line.width = Pt(lw)
    ser.smooth = True
    ser.marker.style = XL_MARKER_STYLE.NONE

    ca = ch.category_axis
    ca.has_major_gridlines = False
    ca.tick_labels.font.size = Pt(fs)
    ca.tick_labels.font.name = FONT
    ca.format.line.color.rgb = _rgb(MUTED)
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _rgb("#E0E0E0")
    va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.color.rgb = _rgb(MUTED)
    va.tick_labels.font.size = Pt(fs)
    va.tick_labels.font.name = FONT
    if ytitle:
        _axis_title(va, ytitle, fs + 1)
    if xtitle:
        _axis_title(ca, xtitle, fs + 1)
    return ch
