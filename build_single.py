#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""نسخهٔ تک‌اسلایدی (یک صفحه) از معماری Prometheus — متن‌ها فارسی و برچسب فازها عمودی."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

from build_pptx import (add_par, arrow, box, text, NAVY, NAVY_SOFT,
                        WHITE, GREY_TXT, BLUE, GREEN, PURPLE, ORANGE, GREY, NAVYF)

OUT = "prometheus-single-page.pptx"
SW, SH = 13.333, 9.5

CL, CR = 0.42, 11.403
CW = CR - CL
NAVY_X, NAVY_W = 11.663, 1.25

SEC_H = 2.6667
SEC_Y = [1.12, 3.8867, 6.6533]

G4 = G3 = 0.22
W4 = (CW - 3 * G4) / 4
W3 = (CW - 2 * G3) / 3
X4 = [CR - W4 - i * (W4 + G4) for i in range(4)]
C4 = [x + W4 / 2 for x in X4]
X3 = [CR - W3 - i * (W3 + G3) for i in range(3)]
C3 = [x + W3 / 2 for x in X3]


def vbox(slide, x, y, w, h, title, subs, name="phase_label"):
    """کادر سرمه‌ای با متن عمودی (از پایین به بالا)."""
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = name
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = NAVY
    sh.line.fill.background()
    sh.adjustments[0] = 0.06
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    bodyPr.set("vert", "vert270")
    add_par(tf, title, 12, True, WHITE, 0, True)
    for sub in subs:
        add_par(tf, sub, 8, False, RGBColor(0xD5, 0xDE, 0xEC), 4)
    return sh


def dline(slide, x1, y1, x2, y2, name="separator"):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.name = name
    cn.line.color.rgb = RGBColor(0xA9, 0xB2, 0xC0)
    cn.line.width = Pt(1)
    cn.line.dash_style = 4
    return cn


def build():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # ---------- title ----------
    text(s, CL, 0.16, CW, 0.55,
         "پرامتئوس: پیش‌بینی فرآیند با درک تأخیر و سه جریان موازی",
         19, True, NAVY, "ctr", "main_title")
    text(s, CL, 0.70, CW, 0.32,
         "معماری از main.py (مدل اصلی) و utils.py (پیاده‌سازی گراف عصبی پیشوند)",
         8.5, False, GREY_TXT, "ctr", "subtitle")

    # ---------- phase labels (vertical, right side) ----------
    labels = [
        ("فاز ۱", ["مشخصات سیستم", "ورودی • وظایف • خروجی"]),
        ("فاز ۲", ["طراحی معماری", "سه جریان موازی"]),
        ("فاز ۳", ["طراحی جزئی", "لایه‌ها و پارامترها"]),
    ]
    for (t, subs), y in zip(labels, SEC_Y):
        vbox(s, NAVY_X, y, NAVY_W, SEC_H, t, subs, f"phase_label_{t[-1]}")

    for y in (3.8367, 6.6034):
        dline(s, CL, y, CR, y)

    # ================= PHASE 1 =================
    y1, y2, h = 1.22, 2.62, 1.17
    top = [
        ("ورودی: لاگ رویداد (XES)", ["پرونده‌ها، فعالیت‌ها، زمان‌ها", "شناسهٔ هر پرونده"], BLUE),
        ("ساخت پیشوند", ["برش هر پرونده به پیشوند", "کمترین طول پیشوند: ۲"], BLUE),
        ("آمار زمانی و موردی", ["زمان انتظار و مدت کل", "زمان باقی‌مانده و نشانهٔ گلوگاه"], GREEN),
        ("اهداف پیش‌بینی", ["فعالیت بعدی و زمان باقی‌مانده", "کاهش تأخیر و تشخیص گلوگاه"], ORANGE),
    ]
    for (t, ls, col), x in zip(top, X4):
        box(s, x, y1, W4, h, t, ls, col, f"p1_top_{t[:6]}", tsize=9.5, bsize=7.5)
    for a, b in ((0, 1), (1, 2), (2, 3)):
        arrow(s, X4[a], y1 + h / 2, X4[b] + W4, y1 + h / 2, NAVY_SOFT, 1.5, "p1_chain")

    bottom = [
        ("ویژگی‌های ترتیبی و زمانی",
         ["فعالیت قبلی، موقعیت، فاصلهٔ زمانی", "زمان تجمعی، نقش، فراوانی"], BLUE),
        ("گراف پیشوند",
         ["گره‌ها: فعالیت + ویژگی زمان/موقعیت", "یال‌ها: ترتیب و تکرار فعالیت‌ها"], PURPLE),
        ("تنظیمات یادگیری",
         ["یادگیری چندوظیفه‌ای با نظارت", "تقسیم داده بر اساس پرونده"], GREEN),
    ]
    for (t, ls, col), x in zip(bottom, X4[1:]):
        box(s, x, y2, W4, h, t, ls, col, f"p1_bot_{t[:6]}", tsize=9.5, bsize=7.5)
    for x in C4[1:]:
        arrow(s, x, y1 + h, x, y2, NAVY_SOFT, 1.5, "p1_branch")

    # ================= PHASE 2 =================
    box(s, C3[1] - 2.8, 3.95, 5.6, 0.52, "بازنمایی ورودی مشترک",
        ["فعالیت + فعالیت قبلی + موقعیت • کدگذاری زمان، نقش و فراوانی"],
        NAVYF, "p2_shared", tsize=9.5, bsize=7.5, title_color=NAVY)

    ys, hs = 4.65, 0.92
    streams = [
        ("جریان ۱ — Transformer", ["کدگذار علّی", "وابستگی‌های دوربرد"], BLUE),
        ("جریان ۲ — TCN", ["پیچش زمانی", "الگوهای محلی چندمقیاسی"], GREEN),
        ("جریان ۳ — GNN گراف پیشوند", ["شبکهٔ عصبی گراف", "روابط میان فعالیت‌های پیشوند"], PURPLE),
    ]
    for (t, ls, col), x in zip(streams, X3):
        box(s, x, ys, W3, hs, t, ls, col, f"p2_stream_{t[:10]}", tsize=9.5, bsize=7.5)

    for x0, x1 in ((7.0, C3[0]), (C3[1], C3[1]), (4.85, C3[2])):
        arrow(s, x0, 4.47, x1, ys, NAVY_SOFT, 1.4, "p2_fan")

    yf, hf = 5.75, 0.76
    box(s, X3[1], yf, CR - X3[1], hf, "ادغام توجه متقابل (Transformer و TCN)",
        ["هر جریان با توجه به خروجی جریان دیگر بازبینی می‌شود"],
        ORANGE, "p2_cross", tsize=9.5, bsize=7.5)
    box(s, X3[2], yf, W3, hf, "ادغام سه‌جریانه",
        ["ترتیب + گراف + تعامل"], ORANGE, "p2_tri", tsize=9.5, bsize=7.5)
    for x in (C3[0], C3[1], C3[2]):
        arrow(s, x, ys + hs, x, yf, NAVY_SOFT, 1.5, "p2_to_fusion")
    arrow(s, X3[1], yf + hf / 2, X3[2] + W3, yf + hf / 2, NAVY, 1.8, "p2_head")

    # ================= PHASE 3 =================
    rows = [6.68, 7.20, 7.72, 8.24, 8.76]
    h3 = 0.38
    a = [
        ("کدگذاری ورودی (۲۵۶)", ["تعبیهٔ فعالیت، موقعیت و موقعیت نسبی"], BLUE),
        ("کدگذاری زمانی", ["log(Δt) و log(زمان تجمعی)"], BLUE),
        ("نرمال‌سازی و تنظیم", ["LayerNorm • Dropout ۰٫۰۵ • ماسک‌ها"], GREY),
    ]
    b = [
        ("کدگذار Transformer", ["۶ لایه • ۸ سر • پیش‌خور ۱۰۲۴"], BLUE),
        ("بلوک TCN", ["۱۰ پیچش • کرنل ۳ • گشودگی ۱، ۲، ۴"], GREEN),
        ("گراف پیشوند (GNN)", ["۲ لایهٔ GCN + میانگین ماسک‌شده"], PURPLE),
    ]
    for (t, ls, col), x in zip(a, X3):
        box(s, x, rows[0], W3, h3, t, ls, col, f"p3_a_{t[:8]}", tsize=8.5, bsize=6.8)
    for (t, ls, col), x in zip(b, X3):
        box(s, x, rows[1], W3, h3, t, ls, col, f"p3_b_{t[:8]}", tsize=8.5, bsize=6.8)
    for x in C3:
        arrow(s, x, rows[0] + h3, x, rows[1], NAVY_SOFT, 1.3, "p3_ab")

    box(s, X3[1], rows[2], CR - X3[1], h3, "ادغام توجه متقابل",
        ["توجه چندسری + اتصال باقیمانده + FFN"], ORANGE, "p3_cross",
        tsize=8.5, bsize=6.8)
    box(s, X3[2], rows[2], W3, h3, "ادغام ترتیبی",
        ["آخرین موقعیت معتبر پیشوند"], GREY, "p3_pool", tsize=8.5, bsize=6.8)
    for x in C3:
        arrow(s, x, rows[1] + h3, x, rows[2], NAVY_SOFT, 1.3, "p3_bc")

    box(s, CL, rows[3], CW, h3, "ادغام نهایی — MLP سه‌جریانه",
        ["ترتیب + گراف + تعامل • ابعاد ۷۶۸ به ۵۱۲ به ۲۵۶ • بازنمایی مشترک: LayerNorm + GELU"],
        ORANGE, "p3_mlp", tsize=8.5, bsize=6.8)
    for x in C3:
        arrow(s, x, rows[2] + h3, x, rows[3], NAVY_SOFT, 1.3, "p3_cd")

    heads = [
        ("سر فعالیت", ["فعالیت بعدی"], BLUE),
        ("سر زمان باقی‌مانده", ["رگرسیون زمان"], GREEN),
        ("سر بهبود", ["کاهش تأخیر + Tanh"], GREEN),
        ("سر گلوگاه", ["گلوگاه هر موقعیت"], PURPLE),
    ]
    for (t, ls, col), x in zip(heads, X4):
        box(s, x, rows[4], W4, h3, t, ls, col, f"p3_head_{t[:8]}", tsize=8.5, bsize=6.8)
    for x in C4:
        arrow(s, x, rows[3] + h3, x, rows[4], NAVY_SOFT, 1.3, "p3_de")

    text(s, CL, 9.20, CW, 0.22,
         "آموزش: زیان وزنی چندوظیفه‌ای (Cross-Entropy + Huber + Huber + BCE) • AdamW • Warmup + Cosine • برش گرادیان • SWA",
         7, False, GREY_TXT, "ctr", "train_note")

    prs.save(OUT)
    print("saved", OUT, "slides:", len(prs.slides))


if __name__ == "__main__":
    build()
